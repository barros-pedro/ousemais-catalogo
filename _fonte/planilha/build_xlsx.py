import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

d=json.load(open('/tmp/claude-0/-home-claude/ff240896-dbcc-5a99-bec2-22c34b27ffc7/scratchpad/data.json'))
P={p['id']:p for p in d['PRODUCTS']}
CATS=[c for c in d['CATS'] if c['id']!='all']
catlabel={c['id']:c['label'] for c in CATS}

F='Arial'
H_FILL=PatternFill('solid',fgColor='5B0F7A'); H_FONT=Font(name=F,bold=True,color='FFFFFF',size=10)
IN_FILL=PatternFill('solid',fgColor='FFF2CC')   # amarelo claro = preencher
ID_FILL=PatternFill('solid',fgColor='F2F2F2'); ID_FONT=Font(name=F,size=9,color='808080')
BODY=Font(name=F,size=10); BOLD=Font(name=F,size=10,bold=True)
thin=Side(style='thin',color='E5D5E1'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
BRL='R$ #,##0.00;-R$ #,##0.00;"-"'
WRAP=Alignment(wrap_text=True,vertical='top')

def header(ws,cols,row=1):
    for i,(t,w) in enumerate(cols,1):
        c=ws.cell(row=row,column=i,value=t); c.fill=H_FILL; c.font=H_FONT
        c.alignment=Alignment(wrap_text=True,vertical='center'); c.border=BOX
        ws.column_dimensions[c.column_letter].width=w
    ws.row_dimensions[row].height=32

def style_row(ws,r,ncol,inputs,ids=(1,),money=()):
    for c in range(1,ncol+1):
        cell=ws.cell(row=r,column=c); cell.border=BOX
        cell.font=BODY; cell.alignment=WRAP
        if c in ids: cell.fill=ID_FILL; cell.font=ID_FONT
        if c in inputs: cell.fill=IN_FILL
        if c in money: cell.number_format=BRL

wb=Workbook()

# ---------- Como preencher ----------
ws=wb.active; ws.title='Como preencher'
ws.column_dimensions['A'].width=3; ws.column_dimensions['B'].width=95
lines=[
 ('Planilha de produtos — Catálogo Ouse Mais', Font(name=F,size=16,bold=True,color='5B0F7A')),
 ('Preencha os campos em amarelo. Depois me envie a planilha que eu atualizo o site com os valores.', Font(name=F,size=11)),
 ('',BODY),
 ('Legenda',BOLD),
 ('AMARELO = você preenche',BODY),
 ('CINZA = código interno do site, não altere',BODY),
 ('Sem cor = já preenchido com o que está no site (pode corrigir se algo estiver errado)',BODY),
 ('',BODY),
 ('Abas',BOLD),
 ('Produtos: preço, preço antigo (para mostrar "de/por"), ranking de mais vendidos e se está disponível ou esgotado.',BODY),
 ('Kits: preço de cada kit. A planilha soma o preço dos itens separados e calcula quanto o cliente economiza.',BODY),
 ('Categorias: quantos produtos cada categoria tem hoje (calculado automaticamente).',BODY),
 ('Entrega e pagamento: frete, prazo e formas de pagamento aceitas (o pagamento é feito no WhatsApp).',BODY),
 ('',BODY),
 ('Como preencher (exemplo)',BOLD),
 ('Preço: só o número, ex.: 29,90. Deixe vazio se ainda for "sob consulta".',BODY),
 ('Preço antigo: só se o produto estiver em promoção, ex.: 39,90. O site mostra "de R$ 39,90 por R$ 29,90".',BODY),
 ('Ranking mais vendidos: 1 para o que mais vende, 2 para o segundo e assim por diante. Os 3 primeiros ganham o selo "Mais vendido".',BODY),
 ('Disponibilidade: escolha Disponível ou Esgotado na lista.',BODY),
]
for i,(t,f) in enumerate(lines,1):
    c=ws.cell(row=i,column=2,value=t); c.font=f; c.alignment=Alignment(wrap_text=True,vertical='top')
for r in (5,): ws.cell(row=r,column=2).fill=IN_FILL
ws.cell(row=6,column=2).fill=ID_FILL
ws.sheet_view.showGridLines=False

# ---------- Produtos ----------
wp=wb.create_sheet('Produtos')
cols=[('Código do site',18),('Categoria',15),('Produto',30),('Marca',15),('Detalhes',22),('Opções (cores/frases)',34),
      ('Preço (R$)',13),('Preço antigo "de" (R$)',14),('Ranking mais vendidos',12),('Disponibilidade',14),('Observações',30)]
header(wp,cols)
order={c['id']:i for i,c in enumerate(CATS)}
prods=[p for p in d['PRODUCTS'] if not p.get('items')]
prods.sort(key=lambda p:order[p['cat']])
for r,p in enumerate(prods,2):
    opts=''
    if p.get('variants'):
        v=p['variants']; opts=f"{v['label']}: "+', '.join(o['name'].title() if v['type']=='select' else o['name'] for o in v['options'])
        if v.get('note'): opts+=f" ({v['note']})"
    vals=[p['id'],catlabel[p['cat']],p['name'],p['brand'],', '.join(p.get('specs',[])),opts,None,None,None,'Disponível','']
    for c,v in enumerate(vals,1): wp.cell(row=r,column=c,value=v)
    style_row(wp,r,len(cols),inputs=(7,8,9,10,11),money=(7,8))
last_p=1+len(prods)
wp.freeze_panes='C2'
dv=DataValidation(type='list',formula1='"Disponível,Esgotado"',allow_blank=True); wp.add_data_validation(dv); dv.add(f'J2:J{last_p+40}')
dvr=DataValidation(type='whole',operator='between',formula1='1',formula2='999',allow_blank=True,error='Use um número inteiro: 1 para o que mais vende.'); wp.add_data_validation(dvr); dvr.add(f'I2:I{last_p+40}')
dvc=DataValidation(type='list',formula1=f"=Categorias!$B$2:$B${1+len(CATS)}",allow_blank=True); wp.add_data_validation(dvc); dvc.add(f'B2:B{last_p+40}')
wp.cell(row=1,column=8).comment=Comment('Opcional. Preencha só se o produto estiver em promoção; o site mostra "de / por".','Ouse Mais')
wp.cell(row=1,column=7).comment=Comment('Deixe vazio para continuar como "Preço sob consulta".','Ouse Mais')
# espaço para novos produtos
for r in range(last_p+1,last_p+11):
    style_row(wp,r,len(cols),inputs=(2,3,4,5,6,7,8,9,10,11),money=(7,8))
wp.cell(row=last_p+1,column=1,value='(novo)').font=ID_FONT
nr=last_p+12
wp.cell(row=nr,column=1,value='Linhas em amarelo vazias: use para cadastrar produtos novos (mande as fotos junto).').font=Font(name=F,size=9,italic=True,color='808080')

# ---------- Kits ----------
wk=wb.create_sheet('Kits')
cols=[('Código do site',16),('Kit',24),('Item 1',26),('Item 2',26),('Item 3',26),('Soma dos itens separados (R$)',16),
      ('Preço do kit (R$)',14),('Economia (R$)',13),('Economia (%)',11),('Disponibilidade',14),('Observações',26)]
header(wk,cols)
kits=[p for p in d['PRODUCTS'] if p.get('items')]
rng_name=f"Produtos!$C$2:$C${last_p+10}"; rng_price=f"Produtos!$G$2:$G${last_p+10}"
for r,k in enumerate(kits,2):
    items=[P[i]['name'] for i in k['items']]+['']*(3-len(k['items']))
    vals=[k['id'],k['name'],*items,
          f'=IF(C{r}="",0,SUMIF({rng_name},C{r},{rng_price}))+IF(D{r}="",0,SUMIF({rng_name},D{r},{rng_price}))+IF(E{r}="",0,SUMIF({rng_name},E{r},{rng_price}))',
          None,
          f'=IF(OR(G{r}="",F{r}=0),"",F{r}-G{r})',
          f'=IF(OR(G{r}="",F{r}=0),"",(F{r}-G{r})/F{r})',
          'Disponível','']
    for c,v in enumerate(vals,1): wk.cell(row=r,column=c,value=v)
    style_row(wk,r,len(cols),inputs=(7,10,11),money=(6,7,8))
    wk.cell(row=r,column=9).number_format='0%;-0%;"-"'
last_k=1+len(kits)
dv2=DataValidation(type='list',formula1='"Disponível,Esgotado"',allow_blank=True); wk.add_data_validation(dv2); dv2.add(f'J2:J{last_k}')
dvi=DataValidation(type='list',formula1=f"={rng_name}",allow_blank=True); wk.add_data_validation(dvi); dvi.add(f'C2:E{last_k}')
wk.freeze_panes='C2'
n=last_k+2
wk.cell(row=n,column=2,value='Dica: deixe o preço do kit um pouco abaixo da soma dos itens. O site pode mostrar "economize R$ X", que é o que faz o cliente escolher o kit.').font=Font(name=F,size=9,italic=True,color='808080')
wk.cell(row=1,column=6).comment=Comment('Calculado: soma o preço de cada item na aba Produtos. Fica "-" enquanto os preços não forem preenchidos.','Ouse Mais')

# ---------- Categorias ----------
wc=wb.create_sheet('Categorias')
header(wc,[('Ordem no site',10),('Categoria',18),('Produtos no catálogo',14),('Situação',30)])
for r,c in enumerate(CATS,2):
    wc.cell(row=r,column=1,value=r-1); wc.cell(row=r,column=2,value=c['label'])
    if c['id']=='kit': wc.cell(row=r,column=3,value=f'=COUNTA(Kits!$B$2:$B${last_k})')
    else: wc.cell(row=r,column=3,value=f'=COUNTIF(Produtos!$B$2:$B${last_p+10},B{r})')
    wc.cell(row=r,column=4,value=f'=IF(C{r}=0,"Sem produtos: mostra aviso de WhatsApp","Com produtos")')
    style_row(wc,r,4,inputs=(),ids=())
lc=1+len(CATS)
wc.cell(row=lc+1,column=2,value='Total').font=BOLD
wc.cell(row=lc+1,column=3,value=f'=SUM(C2:C{lc})').font=BOLD

# ---------- Pagamento e frete ----------
wf=wb.create_sheet('Entrega e pagamento')
wf.column_dimensions['A'].width=34; wf.column_dimensions['B'].width=42; wf.column_dimensions['C'].width=50
header(wf,[('Informação',34),('Resposta',42),('Para que serve no site',50)])
rows=[('Formas de pagamento aceitas','', 'Ex.: Pix, cartão, dinheiro na entrega. Aparece nos avisos do site; o pagamento continua sendo feito pelo WhatsApp.'),
 ('Parcela no cartão? Em quantas vezes?','', 'Mostrado nos avisos do site.'),
 ('Frete em Campestre','', 'Ex.: grátis, R$ 5, retirada na loja.'),
 ('Frete para outras cidades','', 'Ex.: valor fixo por cidade, Correios, combinado no WhatsApp.'),
 ('Frete grátis acima de (R$)',None, 'Opcional. Cria a faixa "frete grátis acima de R$ X".'),
 ('Prazo de envio','', 'Ex.: envio no mesmo dia útil.'),
 ('Pode retirar na loja? Endereço','', 'Opcional.'),
 ('Embalagem sigilosa','Sim', 'Já está no site.'),]
for r,(a,b,c) in enumerate(rows,2):
    wf.cell(row=r,column=1,value=a); wf.cell(row=r,column=2,value=b); wf.cell(row=r,column=3,value=c)
    style_row(wf,r,3,inputs=(2,) if a!='Embalagem sigilosa' else (),ids=())
    wf.cell(row=r,column=1).font=BOLD
    wf.cell(row=r,column=3).font=Font(name=F,size=9,color='808080')
wf.cell(row=6,column=2).number_format=BRL
wf.freeze_panes='A2'

wb.save('/home/claude/catalogo/planilha/Ouse_Mais_Produtos.xlsx')
print('saved', last_p, last_k)
